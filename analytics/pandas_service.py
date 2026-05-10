"""
Pandas-сервіс для аналітики, де ORM незручний:

- Resample часових рядів (агрегація по 1хв / 1год / 1день)
- Rolling-window детектор аномалій (плавне виявлення сплесків помилок)
- Pivot-таблиці (URL × статус-код)
- Експорт у Excel з кількома вкладками і вбудованим графіком

ORM залишається для швидких простих агрегацій у дашборді.
Pandas використовується для важчих обчислень і експорту.

Архітектурно це другий бекенд, який ховається за тим самим Facade:
views не знають, чи аналіз робить ORM, чи pandas.
"""
from __future__ import annotations

import io
from typing import Iterable

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows

from .models import LogEntry, LogFile


class PandasAnalyticsService:
    """Розширена аналітика на pandas. Ховається за фасадом."""

    # ------------------------------------------------------------------
    # Завантаження даних у DataFrame (з опціональними фільтрами)
    # ------------------------------------------------------------------
    def to_dataframe(
        self,
        log_file: LogFile,
        filters: dict | None = None,
    ) -> pd.DataFrame:
        """
        Завантажує записи у DataFrame з опціональним застосуванням фільтрів:
            filters = {
                "status_class": "2xx" | "3xx" | "4xx" | "5xx" | "all",
                "method":       "GET" | "POST" | ... | "all",
                "from":         ISO datetime str,
                "to":           ISO datetime str,
                "path_contains": str,
            }
        """
        qs = log_file.entries.values(
            "timestamp", "ip_address", "method", "path",
            "status_code", "response_size", "browser",
            "os_family", "device_type",
        )
        df = pd.DataFrame.from_records(qs)
        if df.empty:
            return df
        df["timestamp"] = pd.to_datetime(df["timestamp"])

        if not filters:
            return df

        if filters.get("status_class") and filters["status_class"] != "all":
            cls = int(filters["status_class"][0]) * 100
            df = df[(df["status_code"] >= cls) & (df["status_code"] < cls + 100)]

        if filters.get("method") and filters["method"] != "all":
            df = df[df["method"] == filters["method"]]

        if filters.get("from"):
            df = df[df["timestamp"] >= pd.to_datetime(filters["from"])]
        if filters.get("to"):
            df = df[df["timestamp"] <= pd.to_datetime(filters["to"])]

        if filters.get("path_contains"):
            df = df[df["path"].str.contains(filters["path_contains"], na=False)]

        return df

    # ------------------------------------------------------------------
    # Універсальна агрегація з вибором групування і фільтрами
    # ------------------------------------------------------------------
    GROUP_BY_FIELDS = {
        "browser": "browser",
        "os_family": "os_family",
        "device_type": "device_type",
        "method": "method",
        "status_class": None,  # обчислюється
        "ip_address": "ip_address",
        "path": "path",
    }

    def aggregate(
        self,
        log_file: LogFile,
        group_by: str = "method",
        filters: dict | None = None,
        top: int = 10,
    ) -> dict:
        """
        Універсальна агрегація: підраховує кількість запитів,
        згрупованих за `group_by`, з опціональними фільтрами.
        Повертає labels/data для Chart.js.
        """
        df = self.to_dataframe(log_file, filters=filters)
        if df.empty:
            return {"labels": [], "data": [], "group_by": group_by}

        if group_by == "status_class":
            df = df.copy()
            df["__grp"] = (df["status_code"] // 100).astype(str) + "xx"
            grouped = df.groupby("__grp").size().sort_index()
        else:
            field = self.GROUP_BY_FIELDS.get(group_by, "method")
            grouped = (
                df[df[field] != ""]
                .groupby(field).size()
                .sort_values(ascending=False)
                .head(top)
            )
        return {
            "labels": [str(x) for x in grouped.index.tolist()],
            "data": [int(x) for x in grouped.values.tolist()],
            "group_by": group_by,
            "filtered_count": int(len(df)),
        }

    # ------------------------------------------------------------------
    # Status over time (stacked area)
    # ------------------------------------------------------------------
    def status_over_time(
        self, log_file: LogFile, freq: str = "1h",
        filters: dict | None = None,
    ) -> dict:
        """
        Стек-агрегація: ряд для кожного класу статусу (2xx/3xx/4xx/5xx)
        у часі. Зручно для виявлення коли саме почалися помилки.
        """
        df = self.to_dataframe(log_file, filters=filters)
        if df.empty:
            return {"labels": [], "series": {}}

        df = df.copy()
        df["status_class"] = (df["status_code"] // 100).astype(str) + "xx"
        pivot = (
            df.set_index("timestamp")
              .groupby([pd.Grouper(freq=freq), "status_class"]).size()
              .unstack(fill_value=0)
              .sort_index()
        )
        # Гарантуємо порядок класів
        for cls in ["2xx", "3xx", "4xx", "5xx"]:
            if cls not in pivot.columns:
                pivot[cls] = 0
        pivot = pivot[["2xx", "3xx", "4xx", "5xx"]]

        return {
            "labels": [t.strftime("%Y-%m-%d %H:%M") for t in pivot.index],
            "series": {
                cls: [int(v) for v in pivot[cls].tolist()]
                for cls in pivot.columns
            },
        }

    # ------------------------------------------------------------------
    # Heatmap день тижня × година
    # ------------------------------------------------------------------
    def hour_weekday_heatmap(
        self, log_file: LogFile, filters: dict | None = None,
    ) -> dict:
        """
        Матриця 7×24 (день тижня × година доби).
        Корисно щоб зрозуміти патерни активності.
        """
        df = self.to_dataframe(log_file, filters=filters)
        if df.empty:
            return {"matrix": [[0] * 24 for _ in range(7)], "max": 0}

        df = df.copy()
        df["dow"] = df["timestamp"].dt.dayofweek  # 0=Mon
        df["hour"] = df["timestamp"].dt.hour
        pivot = (
            df.groupby(["dow", "hour"]).size()
              .unstack(fill_value=0)
              .reindex(index=range(7), columns=range(24), fill_value=0)
        )
        max_val = int(pivot.values.max()) if pivot.size else 0
        return {
            "matrix": pivot.values.astype(int).tolist(),
            "max": max_val,
            "weekdays": ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Нд"],
        }

    # ------------------------------------------------------------------
    # Histogram розміру відповіді
    # ------------------------------------------------------------------
    def response_size_histogram(
        self, log_file: LogFile, bins: int = 12,
        filters: dict | None = None,
    ) -> dict:
        """Гістограма розподілу розміру відповіді (logarithmic bins)."""
        df = self.to_dataframe(log_file, filters=filters)
        if df.empty:
            return {"labels": [], "data": []}

        sizes = df["response_size"]
        sizes = sizes[sizes > 0]  # відкидаємо порожні (404 і т.п.)
        if sizes.empty:
            return {"labels": [], "data": []}

        # Логарифмічні bin-и
        import numpy as np
        max_size = sizes.max()
        edges = np.unique(np.logspace(0, np.log10(max(max_size, 10)), bins + 1).astype(int))
        if len(edges) < 2:
            edges = np.array([0, int(max_size) + 1])
        counts, edges = np.histogram(sizes, bins=edges)
        labels = [
            f"{self._fmt_bytes(edges[i])}–{self._fmt_bytes(edges[i + 1])}"
            for i in range(len(edges) - 1)
        ]
        return {"labels": labels, "data": counts.tolist()}

    @staticmethod
    def _fmt_bytes(n: int) -> str:
        n = int(n)
        if n < 1024:
            return f"{n}B"
        if n < 1024 * 1024:
            return f"{n // 1024}K"
        return f"{n // (1024 * 1024)}M"

    # ------------------------------------------------------------------
    # Top IPs
    # ------------------------------------------------------------------
    def top_ips(
        self, log_file: LogFile, top: int = 10,
        filters: dict | None = None,
    ) -> dict:
        df = self.to_dataframe(log_file, filters=filters)
        if df.empty:
            return {"labels": [], "data": []}
        grouped = (
            df.groupby("ip_address").size()
              .sort_values(ascending=False).head(top)
        )
        return {
            "labels": grouped.index.tolist(),
            "data": [int(x) for x in grouped.values.tolist()],
        }

    # ------------------------------------------------------------------
    # Time-series: resample
    # ------------------------------------------------------------------
    def resample_traffic(
        self, log_file: LogFile, freq: str = "1h",
        filters: dict | None = None,
    ) -> dict:
        """
        Агрегація запитів по часових інтервалах.
        freq: '1min', '5min', '1h', '1D' (формат pandas offset alias).
        """
        df = self.to_dataframe(log_file, filters=filters)
        if df.empty:
            return {"labels": [], "data": []}
        series = df.set_index("timestamp").resample(freq).size()
        return {
            "labels": [t.strftime("%Y-%m-%d %H:%M") for t in series.index],
            "data": series.tolist(),
        }

    # ------------------------------------------------------------------
    # Rolling-window детектор аномалій
    # ------------------------------------------------------------------
    def rolling_error_rate(
        self,
        log_file: LogFile,
        window: str = "10min",
        threshold: float = 0.1,
    ) -> dict:
        """
        Обчислює ковзне співвідношення 5xx/total у вікні `window`.
        Повертає інтервали, де rate перевищує threshold.

        Це покращення проти простого 'середнє по всьому файлу':
        великий короткий сплеск помилок не розчиняється у середньому.
        """
        df = self.to_dataframe(log_file)
        if df.empty:
            return {"timeline": {"labels": [], "data": []}, "spikes": []}

        df = df.set_index("timestamp").sort_index()
        df["is_error"] = (df["status_code"] >= 500).astype(int)

        # Перетворюємо на регулярний 1-хвилинний ряд (для коректного rolling)
        per_min = df.resample("1min").agg(
            total=("status_code", "count"),
            errors=("is_error", "sum"),
        )
        rolling = per_min.rolling(window).sum()
        rolling["rate"] = (rolling["errors"] / rolling["total"]).fillna(0)

        spikes = rolling[rolling["rate"] > threshold]
        return {
            "timeline": {
                "labels": [t.strftime("%Y-%m-%d %H:%M") for t in rolling.index],
                "data": rolling["rate"].round(4).tolist(),
            },
            "spikes": [
                {
                    "time": t.strftime("%Y-%m-%d %H:%M"),
                    "rate": round(row["rate"], 4),
                    "errors": int(row["errors"]),
                    "total": int(row["total"]),
                }
                for t, row in spikes.iterrows()
            ],
            "threshold": threshold,
            "window": window,
        }

    # ------------------------------------------------------------------
    # Pivot: URL × статус-код
    # ------------------------------------------------------------------
    def paths_status_pivot(
        self, log_file: LogFile, top: int = 10
    ) -> pd.DataFrame:
        """
        Зведена таблиця: рядки = топ-N URL, колонки = класи статус-кодів.
        Зручно для аналізу, які URL частіше падають з 4xx/5xx.
        """
        df = self.to_dataframe(log_file)
        if df.empty:
            return pd.DataFrame()

        df["status_class"] = (df["status_code"] // 100).astype(str) + "xx"
        # Беремо топ-N URL
        top_paths = df["path"].value_counts().head(top).index
        sub = df[df["path"].isin(top_paths)]
        pivot = pd.crosstab(sub["path"], sub["status_class"])
        # Сортуємо рядки за загальною кількістю
        pivot["total"] = pivot.sum(axis=1)
        pivot = pivot.sort_values("total", ascending=False)
        return pivot

    # ------------------------------------------------------------------
    # Експорт у Excel
    # ------------------------------------------------------------------
    def export_to_excel(self, log_file: LogFile) -> bytes:
        """
        Генерує Excel-файл з декількома вкладками:
          1) Огляд — ключові метрики
          2) Трафік за годину — таблиця + lineChart
          3) Топ URL × статус — pivot таблиця + barChart
          4) Браузери — розподіл
        Повертає bytes файлу.
        """
        df = self.to_dataframe(log_file)
        wb = Workbook()

        # --- Sheet 1: Огляд ---
        ws = wb.active
        ws.title = "Огляд"
        ws.append(["Метрика", "Значення"])
        if df.empty:
            ws.append(["Записів", 0])
        else:
            ws.append(["Всього запитів", len(df)])
            ws.append(["Унікальних IP", df["ip_address"].nunique()])
            ws.append(["Сумарний трафік (байт)", int(df["response_size"].sum())])
            ws.append(
                ["Середній розмір (байт)", round(df["response_size"].mean(), 2)]
            )
            errors = (df["status_code"] >= 500).sum()
            ws.append(["Помилок 5xx", int(errors)])
            ws.append(
                ["Доля помилок (%)", round(errors / len(df) * 100, 2)]
            )

        # --- Sheet 2: Трафік за годину ---
        ws2 = wb.create_sheet("Трафік за годину")
        if not df.empty:
            hourly = (
                df.set_index("timestamp")
                .resample("1h").size()
                .reset_index(name="requests")
            )
            hourly["timestamp"] = hourly["timestamp"].dt.strftime("%Y-%m-%d %H:%M")
            for row in dataframe_to_rows(hourly, index=False, header=True):
                ws2.append(row)
            # Додаємо line chart
            if len(hourly) > 1:
                chart = LineChart()
                chart.title = "Запити за годину"
                chart.y_axis.title = "Запитів"
                chart.x_axis.title = "Час"
                data_ref = Reference(
                    ws2, min_col=2, min_row=1, max_row=len(hourly) + 1
                )
                cats_ref = Reference(
                    ws2, min_col=1, min_row=2, max_row=len(hourly) + 1
                )
                chart.add_data(data_ref, titles_from_data=True)
                chart.set_categories(cats_ref)
                ws2.add_chart(chart, "D2")

        # --- Sheet 3: URL × статус ---
        ws3 = wb.create_sheet("URL × статус")
        pivot = self.paths_status_pivot(log_file)
        if not pivot.empty:
            pivot_reset = pivot.reset_index()
            for row in dataframe_to_rows(pivot_reset, index=False, header=True):
                ws3.append(row)
            # Bar chart по топ-URL
            chart = BarChart()
            chart.title = "Топ URL за кількістю запитів"
            chart.y_axis.title = "Кількість"
            data_ref = Reference(
                ws3,
                min_col=pivot_reset.columns.get_loc("total") + 1,
                min_row=1,
                max_row=len(pivot_reset) + 1,
            )
            cats_ref = Reference(
                ws3, min_col=1, min_row=2, max_row=len(pivot_reset) + 1
            )
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws3.add_chart(chart, "H2")

        # --- Sheet 4: Браузери ---
        ws4 = wb.create_sheet("Браузери")
        if not df.empty:
            browsers = (
                df[df["browser"] != ""]
                .groupby("browser").size()
                .sort_values(ascending=False)
                .reset_index(name="requests")
            )
            for row in dataframe_to_rows(browsers, index=False, header=True):
                ws4.append(row)

        # Серіалізуємо
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
